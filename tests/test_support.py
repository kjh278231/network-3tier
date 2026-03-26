from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from src.network3tier.loader import load_network_data
from src.network3tier.webapi.serializers import build_input_payload


HEADER_ROW = 5
HEADER_START_COL = 2


def _write_sheet(workbook: Workbook, sheet_name: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title=sheet_name)
    for col_offset, header in enumerate(headers, start=HEADER_START_COL):
        sheet.cell(row=HEADER_ROW, column=col_offset, value=header)
    for row_offset, row_values in enumerate(rows, start=HEADER_ROW + 1):
        for col_offset, value in enumerate(row_values, start=HEADER_START_COL):
            sheet.cell(row=row_offset, column=col_offset, value=value)


def build_workbook(path: Path, scenario: str = "valid") -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    simulation_headers = ["Simulation Name", "Structure", "Warehouse Qty", "Speed (km/h)", "Coverage (hour)"]
    plant_headers = ["Plant ID", "Location Name", "Product Qty", "Shipment Qty", "Latitude", "Longitude"]
    warehouse_headers = [
        "Warehouse ID",
        "Location Name",
        "Capacity Qty",
        "Fixed Cost",
        "Operation Cost",
        "Latitude",
        "Longitude",
        "Active Y/N",
    ]
    customer_headers = [
        "Customer ID",
        "Location Name",
        "Do Qty",
        "Shipment Qty",
        "Latitude",
        "Longitude",
        "Mapping ID",
    ]
    pw_headers = ["Plant ID", "Warehouse ID", "Distance (km)", "Distance Type", "Trns Cost"]
    wc_headers = ["Warehouse ID", "Customer ID", "Distance (km)", "Distance Type", "Trns Cost"]

    plant_qty = 100
    customer_rows = [
        ["C1", "Customer 1", 50, 1, 37.1, 127.1, None],
        ["C2", "Customer 2", 50, 1, 37.2, 127.2, None],
    ]
    if scenario == "demand_exceeds_supply":
        plant_qty = 80

    _write_sheet(
        workbook,
        "simulation",
        simulation_headers,
        [["Scenario A", "3-tier", 2, 60, 8]],
    )
    _write_sheet(
        workbook,
        "plant",
        plant_headers,
        [["P1", "Plant 1", plant_qty, 10, 37.5, 127.0]],
    )
    _write_sheet(
        workbook,
        "warehouse",
        warehouse_headers,
        [
            ["W1", "Warehouse 1", 100, 10, 1, 37.31, 127.31, "Y"],
            ["W2", "Warehouse 2", 100, 10, 1, 37.32, 127.32, "Y"],
            ["W3", "Warehouse 3", 100, 100, 5, 37.33, 127.33, "Y"],
        ],
    )
    _write_sheet(workbook, "customer", customer_headers, customer_rows)
    _write_sheet(
        workbook,
        "plantWarehouseCost",
        pw_headers,
        [
            ["P1", "W1", 100, "ROAD", 1],
            ["P1", "W2", 120, "ROAD", 1],
            ["P1", "W3", 300, "ROAD", 5],
        ],
    )
    _write_sheet(
        workbook,
        "warehouseCustomerCost",
        wc_headers,
        [
            ["W1", "C1", 10, "ROAD", 1],
            ["W1", "C2", 60, "ROAD", 4],
            ["W2", "C1", 55, "ROAD", 4],
            ["W2", "C2", 10, "ROAD", 1],
            ["W3", "C1", 90, "ROAD", 6],
            ["W3", "C2", 90, "ROAD", 6],
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build_input_json(path: Path, scenario: str = "valid") -> Path:
    workbook_path = build_workbook(path.with_suffix(".xlsx"), scenario=scenario)
    data = load_network_data(workbook_path)
    payload = build_input_payload(data)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
